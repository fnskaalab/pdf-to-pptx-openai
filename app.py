import os
import json
from openai import OpenAI
from dotenv import load_dotenv
from pptx import Presentation
from pptx.util import Inches
import requests
from io import BytesIO
from pdfminer.high_level import extract_text

load_dotenv()
API_KEY = os.getenv("OPENAI_API_KEY")
client = OpenAI(api_key=API_KEY)

def lire_pdf(fichier_pdf):
    """Extrait le texte d'un fichier PDF."""
    try:
        texte = extract_text(fichier_pdf)
        return texte.strip()
    except Exception as e:
        print(f"Erreur lors de la lecture du PDF : {e}")
        return None

def generer_resume(texte_pdf):
    """Envoie le texte du PDF à OpenAI pour générer une structure de présentation."""
    prompt = (
        "Tu es un expert en rédaction technique. Voici un document que je souhaite résumer sous forme de présentation PowerPoint.\n"
        "Génère un JSON contenant une liste de slides avec : un titre, un sous-titre, un contenu et une URL d'image pertinente.\n\n"
        f"Document : {texte_pdf}\n\n"
        "Réponds uniquement avec un JSON bien structuré sous cette forme :\n"
        '{ "slides": [ { "titre": "...", "sous_titre": "...", "contenu": "...", "image": "..." } ] }'
    )

    try:
        chat_completion = client.chat.completions.create(
            messages=[{"role": "user", "content": prompt}],
            model="gpt-4",
        )
        reponse_json = chat_completion.choices[0].message.content.strip()
        return json.loads(reponse_json)  # Convertir en dict Python
    except Exception as e:
        print(f"Erreur OpenAI : {e}")
        return None

def creer_pptx(data, nom_fichier="presentation_IA.pptx"):
    """Génère une présentation PowerPoint à partir des données JSON."""
    if not data or "slides" not in data:
        print("Données invalides pour générer le PPTX.")
        return

    prs = Presentation()
    
    for slide_data in data["slides"]:
        slide_layout = prs.slide_layouts[1]  # Mise en page avec titre et contenu
        slide = prs.slides.add_slide(slide_layout)

        title = slide.shapes.title
        subtitle = slide.placeholders[1]

        title.text = slide_data.get("titre", "Titre manquant")
        subtitle.text = f"{slide_data.get('sous_titre', '')}\n\n{slide_data.get('contenu', '')}"

        # Ajout de l'image si disponible
        if "image" in slide_data:
            try:
                response = requests.get(slide_data["image"])
                if response.status_code == 200:
                    image_stream = BytesIO(response.content)
                    left, top, width, height = Inches(1), Inches(3), Inches(8), Inches(4.5)
                    slide.shapes.add_picture(image_stream, left, top, width, height)
            except Exception as e:
                print(f"Erreur lors du téléchargement de l'image : {e}")

    prs.save(nom_fichier)
    print(f"✅ Fichier PPTX généré : {nom_fichier}")

pdf_path = "document.pdf"  # Remplace par ton fichier PDF
texte = lire_pdf(pdf_path)

if texte:
    json_slides = generer_resume(texte)
    if json_slides:
        creer_pptx(json_slides)
